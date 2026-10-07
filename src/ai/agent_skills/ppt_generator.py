import logging
import json
import os
import re
from datetime import datetime
try:
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.enum.text import PP_ALIGN
except Exception:
    Presentation = None
    Inches = None
    Pt = None
    PP_ALIGN = None

logger = logging.getLogger(__name__)

class PPTGenerator:
    """
    Agent Skill: PPT Generation (PPT Generator)
    Converts the finalized draft into a scientific presentation PPTX file.
    """

    def __init__(self, llm_service, model="llama3"):
        self.llm_service = llm_service
        self.model = model
        self.output_dir = "outputs/ppt"
        os.makedirs(self.output_dir, exist_ok=True)

    def generate_ppt(self, draft_text: str, template_style="Standard"):
        """
        Main entry point for PPT generation loop.
        """
        logger.info(f"Generating PPT for draft (Length: {len(draft_text)} chars)...")
        if Presentation is None:
            return {"status": "failed", "error": "python-pptx/lxml 不可用"}
        
        # 1. Structure Analysis (LLM)
        slide_structure = self._parse_to_structure(draft_text)
        if not slide_structure:
            logger.error("Failed to parse draft structure.")
            return {"status": "failed", "error": "Structure parsing failed"}
            
        # 2. PPTX Creation (python-pptx)
        ppt_path = self._create_pptx_file(slide_structure, template_style)
        
        return {
            "status": "success",
            "file_path": ppt_path,
            "slide_count": len(slide_structure)
        }

    def _parse_to_structure(self, text):
        """
        Use LLM to break down text into slides JSON.
        """
        prompt = (
            f"You are a Research Presentation Specialist. Convert the following research paper text into a structured presentation outline JSON.\n"
            f"Requirements:\n"
            f"1. Create 8-12 slides total.\n"
            f"2. Include 'Title', 'Introduction', 'Methods', 'Results', 'Discussion', 'Conclusion' slides.\n"
            f"3. For each slide, provide a 'title' and a list of 3-5 'bullets' (concise points).\n"
            f"4. Output strictly valid JSON array of objects.\n\n"
            f"Draft Text:\n{text[:6000]}...\n\n" # Truncate if too long for context window
            f"Output JSON Format:\n"
            f"[{{\"title\": \"Slide Title\", \"layout\": \"Content\", \"bullets\": [\"Point 1\", \"Point 2\"]}}]"
        )
        
        try:
            # Timeout set to None for large documents
            res = self.llm_service.generate(self.model, prompt, timeout=None)
            json_str = res.get("response", "").strip()
            
            # Extract JSON if wrapped in markdown
            if "```json" in json_str:
                json_str = json_str.split("```json")[1].split("```")[0].strip()
            elif "```" in json_str:
                json_str = json_str.split("```")[1].split("```")[0].strip()
                
            slides = json.loads(json_str)
            return slides if isinstance(slides, list) else None
        except Exception as e:
            logger.error(f"LLM Parsing failed: {e}")
            # Fallback: Simple heuristic splitting
            return self._heuristic_fallback(text)

    def _heuristic_fallback(self, text):
        """
        Simple fallback if LLM fails JSON generation.
        """
        logger.warning("Using heuristic fallback for slides.")
        slides = []
        slides.append({"title": "Research Paper Presentation", "layout": "Title", "bullets": []})
        
        # Simple split by headers
        sections = re.split(r'\n#+\s+', text)
        for sec in sections:
            if not sec.strip(): continue
            lines = sec.strip().split('\n')
            title = lines[0]
            content = lines[1:]
            bullets = [line for line in content if line.strip()][:5] # Take first 5 lines
            if bullets:
                slides.append({"title": title, "layout": "Content", "bullets": bullets})
                
        return slides

    def _create_pptx_file(self, slides, style):
        """
        Render slides into PPTX.
        """
        prs = Presentation()
        
        # Define layouts (0: Title, 1: Title+Content)
        # Note: Layout indices depend on the master slide. Standard usually:
        # 0: Title Slide
        # 1: Title and Content
        
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        filename = f"Research_Presentation_{timestamp}.pptx"
        filepath = os.path.join(self.output_dir, filename)
        
        for slide_data in slides:
            layout_idx = 0 if slide_data.get("layout") == "Title" else 1
            slide_layout = prs.slide_layouts[layout_idx]
            slide = prs.slides.add_slide(slide_layout)
            
            # Set Title
            if slide.shapes.title:
                slide.shapes.title.text = slide_data.get("title", "Untitled")
                
            # Set Content
            if layout_idx == 1 and slide.placeholders[1]:
                tf = slide.placeholders[1].text_frame
                tf.clear() # Clear default
                
                for bullet in slide_data.get("bullets", []):
                    p = tf.add_paragraph()
                    p.text = str(bullet)
                    p.level = 0
                    
        prs.save(filepath)
        logger.info(f"PPT saved to {filepath}")
        return filepath
