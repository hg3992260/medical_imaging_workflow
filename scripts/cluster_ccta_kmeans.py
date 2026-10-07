import os
import sys
import argparse
import csv
from typing import List, Tuple, Optional
import numpy as np
import matplotlib.pyplot as plt
from scipy.cluster.vq import kmeans2
import openpyxl

FEATURES = [
    "HyperQRS Amax",
    "HyperQRS xRMS80",
    "HyperQRS RMS0.1Am",
    "HyperQRS xKurt",
    "HyperQRS xRaz",
]
LABEL = "CCTA_result"

def find_dataset_path(explicit_path: Optional[str]) -> Optional[str]:
    if explicit_path:
        return explicit_path if os.path.isfile(explicit_path) else None
    search_root = os.path.abspath(os.path.join(os.path.dirname(__file__), os.pardir, os.pardir))
    candidates = []
    for dirpath, dirnames, filenames in os.walk(search_root):
        for name in filenames:
            if name.lower().endswith((".xlsx", ".csv")):
                candidates.append(os.path.join(dirpath, name))
    for p in candidates:
        if has_required_columns(p):
            return p
    return None

def has_required_columns(path: str) -> bool:
    try:
        headers = get_headers(path)
        return headers is not None and all(h in headers for h in FEATURES + [LABEL])
    except Exception:
        return False

def get_headers(path: str) -> Optional[List[str]]:
    ext = os.path.splitext(path)[1].lower()
    if ext == ".csv":
        with open(path, "r", encoding="utf-8", newline="") as f:
            reader = csv.reader(f)
            for row in reader:
                if row:
                    return [c.strip() for c in row]
        return None
    if ext == ".xlsx":
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        ws = wb.active
        first_row = next(ws.iter_rows(min_row=1, max_row=1, values_only=True))
        wb.close()
        if first_row:
            return [str(c).strip() if c is not None else "" for c in first_row]
        return None
    return None

def read_table(path: str) -> Tuple[np.ndarray, dict]:
    ext = os.path.splitext(path)[1].lower()
    if ext == ".csv":
        with open(path, "r", encoding="utf-8", newline="") as f:
            reader = csv.DictReader(f)
            rows = list(reader)
    elif ext == ".xlsx":
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        ws = wb.active
        rows = []
        headers = get_headers(path)
        for r in ws.iter_rows(min_row=2, values_only=True):
            row = {}
            for i, h in enumerate(headers):
                row[h] = r[i] if i < len(r) else None
            rows.append(row)
        wb.close()
    else:
        rows = []
    data = {}
    for key in FEATURES + [LABEL]:
        vals = []
        for row in rows:
            v = row.get(key)
            try:
                vals.append(float(v))
            except Exception:
                vals.append(np.nan)
        data[key] = np.array(vals, dtype=float)
    mask = np.ones(len(rows), dtype=bool)
    for key in FEATURES + [LABEL]:
        mask &= ~np.isnan(data[key])
    for key in data:
        data[key] = data[key][mask]
    arr = np.stack([data[k] for k in FEATURES + [LABEL]], axis=1)
    return arr, data

def synthetic_data(n: int = 200) -> Tuple[np.ndarray, dict]:
    rng = np.random.default_rng(42)
    x1 = rng.normal(0, 1, n)
    x2 = rng.normal(0, 1, n) * 0.5 + 0.5
    x3 = rng.normal(0, 1, n) * 1.5
    x4 = rng.normal(0, 1, n) + (rng.random(n) > 0.5).astype(float)
    x5 = rng.normal(0, 1, n) * 0.8
    y = ((x1 + x2 + x4) > 1.0).astype(float)
    data = {
        FEATURES[0]: x1,
        FEATURES[1]: x2,
        FEATURES[2]: x3,
        FEATURES[3]: x4,
        FEATURES[4]: x5,
        LABEL: y,
    }
    arr = np.stack([data[k] for k in FEATURES + [LABEL]], axis=1)
    return arr, data

def kmeans_clusters(values: np.ndarray, k: int = 2) -> np.ndarray:
    v = values.astype(float)
    s = np.std(v)
    m = np.mean(v)
    if s == 0:
        s = 1.0
    z = (v - m) / s
    centers, labels = kmeans2(z.reshape(-1, 1), k, minit="++", iter=50)
    return labels

def plot_feature(values: np.ndarray, labels: np.ndarray, target: np.ndarray, name: str, out_dir: str) -> str:
    idx = np.argsort(values)
    v = values[idx]
    c = labels[idx]
    t = target[idx]
    fig = plt.figure(figsize=(8, 4))
    plt.scatter(np.arange(len(v)), v, c=c, cmap="tab10", s=30)
    plt.scatter(np.arange(len(v)), v, c=(t > 0).astype(int), cmap="Greys", s=10, marker="x")
    plt.title(name)
    plt.xlabel("index")
    plt.ylabel("value")
    out_path = os.path.join(out_dir, name.replace(" ", "_") + ".png")
    os.makedirs(os.path.dirname(out_path), exist_ok=True)
    plt.tight_layout()
    plt.savefig(out_path)
    plt.close(fig)
    return out_path

def run(path: Optional[str]) -> List[str]:
    dataset_path = find_dataset_path(path)
    if dataset_path is None:
        arr, data = synthetic_data()
    else:
        arr, data = read_table(dataset_path)
    outputs = []
    out_dir = os.path.abspath(os.path.join(os.path.dirname(__file__), os.pardir, "outputs", "cluster_plots"))
    target = data[LABEL]
    for name in FEATURES:
        vals = data[name]
        cls = kmeans_clusters(vals, 2)
        p = plot_feature(vals, cls, target, name, out_dir)
        outputs.append(p)
    return outputs

def main(argv: Optional[List[str]] = None) -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--path", type=str, default=None)
    args = parser.parse_args(argv)
    outs = run(args.path)
    for p in outs:
        print(p)
    return 0

if __name__ == "__main__":
    sys.exit(main())

